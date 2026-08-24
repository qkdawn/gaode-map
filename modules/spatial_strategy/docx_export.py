from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


_IMAGE_RE = re.compile(r"^!\[([^]]*)\]\(([^)]+)\)$")
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.+?)\s*$")
_NUMBER_RE = re.compile(r"^\s*(\d+)[.)]\s+(.+?)\s*$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _set_font(style, name: str = "Calibri", size: int = 11, *, color: str = "000000") -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend((field_begin, instruction, field_end))


def _set_document_zoom(document: Document, percent: int = 100) -> None:
    settings = document.settings.element
    zoom = settings.find(qn("w:zoom"))
    if zoom is None:
        zoom = OxmlElement("w:zoom")
        settings.insert(0, zoom)
    zoom.set(qn("w:val"), "bestFit")
    zoom.set(qn("w:percent"), str(percent))


def _add_inline_markdown(paragraph, text: str) -> None:
    cursor = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def _new_numbering_id(document: Document, *, start: int) -> int:
    numbering = document.part.numbering_part.element
    style_num_id = document.styles["List Number"]._element.pPr.numPr.numId.val
    style_numbering = next(
        item for item in numbering.findall(qn("w:num"))
        if int(item.get(qn("w:numId"))) == int(style_num_id)
    )
    abstract_num_id = style_numbering.find(qn("w:abstractNumId")).get(qn("w:val"))
    num_id = max(int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))) + 1
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    number.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    level_override.append(start_override)
    number.append(level_override)
    numbering.append(number)
    return num_id


def _add_numbered_paragraph(document: Document, text: str, *, num_id: int) -> None:
    paragraph = document.add_paragraph(style="List Number")
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id
    _add_inline_markdown(paragraph, text)


def _set_table_row_pagination(row, *, repeat_header: bool = False) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        properties.append(header)


def _set_cell_margins(cell, *, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _table_widths(rows: list[list[str]], total: int = 9360) -> list[int]:
    width = max(len(row) for row in rows)
    weights = []
    for column in range(width):
        longest = max((len(row[column].strip()) if column < len(row) else 0 for row in rows), default=1)
        weights.append(max(10, min(longest, 42)))
    minimum = 1200
    remaining = total - minimum * width
    weight_total = sum(weights)
    result = [minimum + int(remaining * weight / weight_total) for weight in weights]
    result[-1] += total - sum(result)
    return result


def _apply_table_geometry(table, widths: list[int]) -> None:
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(widths[index]))
            tc_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    _set_table_row_pagination(table.rows[0], repeat_header=True)
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.text = ""
        _add_inline_markdown(cell.paragraphs[0], value.strip())
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows[2:] if len(rows) > 1 and all(set(cell.strip()) <= {"-", ":", " "} for cell in rows[1]) else rows[1:]:
        cells = table.add_row().cells
        _set_table_row_pagination(table.rows[-1])
        for index, value in enumerate(row):
            cells[index].text = ""
            _add_inline_markdown(cells[index].paragraphs[0], value.strip())
    _apply_table_geometry(table, _table_widths(rows))


def _resolve_image(base_dir: Path, raw_path: str) -> Path | None:
    path = Path(raw_path.strip())
    candidate = (base_dir / path).resolve() if not path.is_absolute() else path.resolve()
    return candidate if candidate.is_file() else None


def write_markdown_docx(markdown_path: Path, output_path: Path, *, title: str = "") -> Path:
    """Render the report markdown into a readable Word document."""
    markdown = markdown_path.read_text(encoding="utf-8")
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _set_document_zoom(document)
    normal = document.styles["Normal"]
    _set_font(normal, size=11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    style_tokens = {
        "Title": (22, "0B2545", 0, 16),
        "Heading 1": (16, "2E74B5", 12, 6),
        "Heading 2": (13, "2E74B5", 10, 5),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for style_name, (size, color, before, after) in style_tokens.items():
        style = document.styles[style_name]
        _set_font(style, size=size, color=color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for list_style_name in ("List Bullet", "List Number"):
        list_style = document.styles[list_style_name]
        _set_font(list_style, size=11)
        list_style.paragraph_format.left_indent = Inches(0.5)
        list_style.paragraph_format.first_line_indent = Inches(-0.25)
        list_style.paragraph_format.space_after = Pt(8)
        list_style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("城市更新空间策略")
    header_run.font.name = "Calibri"
    header_rpr = header_run._element.get_or_add_rPr()
    header_fonts = header_rpr.rFonts
    if header_fonts is None:
        header_fonts = OxmlElement("w:rFonts")
        header_rpr.append(header_fonts)
    header_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    header_run.font.size = Pt(8.5)
    header_run.font.color.rgb = RGBColor.from_string("68727D")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("第 ")
    footer_run.font.size = Pt(8.5)
    footer_run.font.color.rgb = RGBColor.from_string("68727D")
    _add_page_number(footer)
    footer_end = footer.add_run(" 页")
    footer_end.font.size = Pt(8.5)
    footer_end.font.color.rgb = RGBColor.from_string("68727D")

    lines = markdown.splitlines()
    index = 0
    active_numbering_id: int | None = None
    expected_number: int | None = None
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        image_match = _IMAGE_RE.match(line)
        if image_match:
            active_numbering_id = expected_number = None
            image_path = _resolve_image(markdown_path.parent, image_match.group(2))
            if image_path:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Inches(6.1))
                caption = document.add_paragraph(image_match.group(1))
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.style = "Caption"
            index += 1
            continue
        heading_match = _HEADING_RE.match(line)
        if heading_match:
            active_numbering_id = expected_number = None
            level = min(len(heading_match.group(1)), 3)
            text = heading_match.group(2).strip()
            if level == 1 and not document.paragraphs:
                document.add_paragraph(text, style="Title")
            else:
                document.add_paragraph(text, style=f"Heading {level}")
            index += 1
            continue
        if line.startswith("|"):
            active_numbering_id = expected_number = None
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            _add_table(document, rows)
            continue
        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            active_numbering_id = expected_number = None
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_markdown(paragraph, bullet_match.group(1))
            index += 1
            continue
        number_match = _NUMBER_RE.match(line)
        if number_match:
            number = int(number_match.group(1))
            if active_numbering_id is None or expected_number != number:
                active_numbering_id = _new_numbering_id(document, start=number)
            _add_numbered_paragraph(document, number_match.group(2), num_id=active_numbering_id)
            expected_number = number + 1
            index += 1
            continue
        if line in {"---", "***"}:
            active_numbering_id = expected_number = None
            index += 1
            continue
        active_numbering_id = expected_number = None
        paragraph = document.add_paragraph()
        _add_inline_markdown(paragraph, line)
        index += 1

    if title and not document.paragraphs:
        document.add_paragraph(title, style="Title")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
