from __future__ import annotations

import asyncio
from unittest.mock import AsyncMock
from uuid import UUID
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from docx.oxml.ns import qn

from modules.spatial_strategy.docx_export import write_markdown_docx
from modules.spatial_strategy.reporting import (
    FeishuConfig,
    FeishuReportSender,
    STRATEGY_CHAPTERS,
    SpatialStrategyReportStore,
    build_spatial_strategy_report,
    compose_spatial_strategy_report,
)
from modules.spatial_strategy.schemas import SpatialStrategyReportFinalizeRequest


RUN_ID = UUID("31d81c16-1bfc-4bf6-a9e8-cfa36d0a4df2")


def _chapters() -> list[dict]:
    chapters = []
    for unit_id, title in STRATEGY_CHAPTERS:
        content = f"{title}形成明确判断。\n\n### 行动建议\n\n落实与本章判断对应的首期动作。"
        if unit_id == "positioning":
            content = "项目应以社区日常使用为底盘，以文化协作为增量。\n\n### 候选比较\n\n公共文化院落优于普通商业街区。"
        citations = []
        if unit_id == "project_basis":
            citations = [{
                "citation_id": "project:source-1",
                "title": "项目材料原文",
                "source_type": "project_document",
                "source_locator": "第1页",
                "record_ref": None,
            }]
        if unit_id == "regional_role":
            citations = [
                {
                    "citation_id": "project:source-1",
                    "title": "项目材料原文",
                    "source_type": "project_document",
                    "source_locator": "第1页",
                    "record_ref": None,
                },
                {
                    "citation_id": "spatial:source-2",
                    "title": "空间查询",
                    "source_type": "project_computed_result",
                    "source_locator": "结果2",
                    "record_ref": "current:dataset:poi/2",
                },
            ]
        chapters.append({"unit_id": unit_id, "title": title, "content": content, "citations": citations})
    return chapters


def _request() -> SpatialStrategyReportFinalizeRequest:
    return SpatialStrategyReportFinalizeRequest(
        run_id=RUN_ID,
        history_id="history-1",
        project_question="判断项目空间策略",
        project_context={"project": {"name": "测试项目"}},
        chapters=_chapters(),
        visual_assets=[
            {
                "kind": "table",
                "section_id": "regional_role",
                "title": "区域角色比较",
                "markdown": "| 选项 | 结论 |\n|---|---|\n| 连接节点 | 推荐 |",
            },
            {
                "kind": "table",
                "section_id": "audience_use",
                "title": "客群任务",
                "markdown": "| 客群 | 任务 |\n|---|---|\n| 居民 | 日常使用 |",
            },
            {
                "kind": "table",
                "section_id": "phasing",
                "title": "首期行动顺序",
                "markdown": "| 顺序 | 行动 |\n|---:|---|\n| 1 | 改善连接 |",
            },
        ],
    )


def test_report_deterministically_renders_eleven_chapters_and_citations():
    report = build_spatial_strategy_report(_request())

    assert report["title"] == "测试项目空间策略与行动方案"
    assert sum(line.startswith("## ") for line in report["markdown"].splitlines()) == 13
    assert "## 4. 客群与使用" in report["markdown"]
    assert report["markdown"].index("## 1. 项目材料与项目基础") < report["markdown"].index("## 11. 首期闭环与后续分期")
    assert report["summary"] == "项目应以社区日常使用为底盘，以文化协作为增量。"
    assert [citation["citation_id"] for citation in report["citations"]] == [
        "project:source-1",
        "spatial:source-2",
    ]
    assert "## 参考来源" in report["markdown"]
    assert "- [1] 项目材料原文；第1页" in report["markdown"]
    assert [chapter["unit_id"] for chapter in report["chapters"]] == [item[0] for item in STRATEGY_CHAPTERS]


def test_report_rewrites_chapter_local_citations_to_global_labels():
    request = _request().model_copy(deep=True)
    request.chapters[0].content += "\n\n项目材料支持该判断。[C1]"
    request.chapters[0].citations[0]["citation_id"] = "C1"
    request.chapters[1].content += "\n\n同一材料与空间结果共同支持该判断。[C2][S1]"
    request.chapters[1].citations[0]["citation_id"] = "C2"
    request.chapters[1].citations[1]["citation_id"] = "S1"

    report = build_spatial_strategy_report(request)

    assert "[C1]" not in report["markdown"]
    assert "[C2]" not in report["markdown"]
    assert "[S1]" not in report["markdown"]
    assert "项目材料支持该判断。[1]" in report["markdown"]
    assert "同一材料与空间结果共同支持该判断。[1][2]" in report["markdown"]
    assert "[1]" in report["chapters"][0]["content"]
    assert report["markdown"].count("- [1] 项目材料原文；第1页") == 1


def test_report_rewrites_fullwidth_chapter_citations_to_global_labels():
    request = _request()
    request.chapters[0].content += "\n\n全角引用也应被转换。【C1】"
    request.chapters[0].citations[0]["citation_id"] = "C1"

    report = build_spatial_strategy_report(request)

    assert "【C1】" not in report["markdown"]
    assert "全角引用也应被转换。[1]" in report["markdown"]


def test_report_rejects_unmapped_chapter_citation():
    request = _request().model_copy(deep=True)
    request.chapters[0].content += "\n\n缺少来源映射。[C9]"

    with pytest.raises(ValueError, match="chapter_citation_mapping_missing:C9"):
        build_spatial_strategy_report(request)


def test_report_rejects_missing_or_reordered_chapters():
    missing = _request().model_copy(deep=True)
    missing.chapters.pop()
    with pytest.raises(ValueError, match="report_requires_ordered_strategy_chapters"):
        build_spatial_strategy_report(missing)

    reordered = _request().model_copy(deep=True)
    reordered.chapters[0], reordered.chapters[1] = reordered.chapters[1], reordered.chapters[0]
    with pytest.raises(ValueError, match="report_requires_ordered_strategy_chapters"):
        build_spatial_strategy_report(reordered)


def test_report_requires_three_to_five_visuals():
    request = _request().model_copy(update={"visual_assets": []})
    with pytest.raises(ValueError, match="report_requires_three_to_five_visuals"):
        build_spatial_strategy_report(request)


def test_report_removes_duplicate_chapter_heading_and_normalizes_inner_headings():
    request = _request().model_copy(deep=True)
    request.chapters[0].content = "## 项目材料与项目基础\n\n## 比较\n\n材料形成明确边界。"

    report = build_spatial_strategy_report(request)

    assert report["markdown"].count("## 1. 项目材料与项目基础") == 1
    assert "### 比较" in report["markdown"]
    assert "\n## 比较\n" not in report["markdown"]


def test_report_rejects_internal_workflow_language_but_allows_g_roads():
    request = _request().model_copy(deep=True)
    request.chapters[1].content = "本章读取 decision_state 后形成结论。"
    with pytest.raises(ValueError, match="contains_internal_terms"):
        build_spatial_strategy_report(request)

    request = _request().model_copy(deep=True)
    request.chapters[1].content = "项目沿 G318 国道形成东西向联系，应组织慢行接驳。"
    assert "G318 国道" in build_spatial_strategy_report(request)["markdown"]

    request = _request().model_copy(deep=True)
    request.chapters[6].content = "居民、运营方和内容伙伴共同形成日常工作流。"
    assert "共同形成日常工作流" in build_spatial_strategy_report(request)["markdown"]


def test_report_uses_project_question_when_saved_scope_name_is_generic():
    request = _request().model_copy(deep=True)
    request.project_context["project"]["name"] = "15min - 112.9863,28.2208 - 2635 POIs"
    request.project_question = "完成长沙县人民政府原址城市更新项目的综合空间分析。"

    report = build_spatial_strategy_report(request)

    assert report["title"] == "长沙县人民政府原址城市更新项目空间策略与行动方案"


def test_report_store_writes_run_scoped_markdown(tmp_path):
    store = SpatialStrategyReportStore(tmp_path)
    path = store.write({"run_id": str(RUN_ID), "markdown": "# 报告\n"})
    assert path == tmp_path / str(RUN_ID) / "spatial-strategy-report.md"
    assert path.read_text(encoding="utf-8") == "# 报告\n"


def test_markdown_report_can_be_rendered_as_docx(tmp_path):
    markdown = tmp_path / "report.md"
    markdown.write_text(
        "# 报告\n\n## 判断\n\n项目材料支持**该方向**。\n\n"
        "1. **条件一**：先核验。\n2. 条件二：再测试。\n\n"
        "新的清单：\n\n1. 新条件一。\n\n"
        "| 类型 | 行动 |\n|---|---|\n| 公共服务 | 先测试 |\n",
        encoding="utf-8",
    )
    output = write_markdown_docx(markdown, tmp_path / "report.docx")

    assert output.is_file() and output.stat().st_size > 1000
    with ZipFile(output) as archive:
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
    assert 'w:val="bestFit"' in settings_xml
    document = Document(output)
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text]
    assert any(run.text == "该方向" and run.bold for paragraph in paragraphs for run in paragraph.runs)
    numbered = [paragraph for paragraph in paragraphs if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None]
    assert len(numbered) == 3
    table = document.tables[0]
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


def test_compose_report_writes_markdown_and_docx(tmp_path):
    result = asyncio.run(compose_spatial_strategy_report(_request(), store=SpatialStrategyReportStore(tmp_path)))

    directory = tmp_path / str(RUN_ID)
    assert (directory / "spatial-strategy-report.md").is_file()
    assert (directory / "spatial-strategy-report.docx").is_file()
    assert result["asset_manifest"]["kind"] == "docx_report"
    assert len(result["chapters"]) == 11


def test_report_embeds_visual_next_to_matching_unit():
    report = build_spatial_strategy_report(_request())
    audience_index = report["markdown"].index("## 4. 客群与使用")
    visual_index = report["markdown"].index("### 客群任务")
    positioning_index = report["markdown"].index("## 5. 地方资源与共同机制")
    assert audience_index < visual_index < positioning_index


def test_feishu_sender_sends_summary_then_word_file(tmp_path):
    markdown = tmp_path / "spatial-strategy-report.md"
    markdown.write_text("# 报告\n", encoding="utf-8")
    document = tmp_path / "spatial-strategy-report.docx"
    document.write_bytes(b"docx")
    response = lambda payload: httpx.Response(200, json=payload, request=httpx.Request("POST", "https://open.feishu.cn/open-apis"))
    client = AsyncMock()
    client.post.side_effect = [
        response({"code": 0, "tenant_access_token": "token"}),
        response({"code": 0, "data": {"message_id": "summary-id"}}),
        response({"code": 0, "data": {"file_key": "file-key"}}),
        response({"code": 0, "data": {"message_id": "file-id"}}),
    ]
    sender = FeishuReportSender(FeishuConfig("app", "secret", "chat", "https://open.feishu.cn/open-apis"), client=client)

    result = asyncio.run(sender.deliver(report={"title": "测试报告", "summary": "结论", "chapters": _chapters()}, markdown_path=markdown, document_path=document))

    assert result["summary_message_id"] == "summary-id"
    assert result["file_message_id"] == "file-id"
    assert client.post.await_count == 4
